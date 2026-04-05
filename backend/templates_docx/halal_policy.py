"""Template: Halal Policy — chính sách Halal tổng thể của doanh nghiệp."""
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from lxml import etree
from . import _base as B

# Defaults when admin hasn't customized
_DEFAULT_META = [
    ("Tiêu chuẩn tham chiếu", "MS 1500:2019, JAKIM Guidelines"),
    ("Bộ phận chịu trách nhiệm", "Ban Quản lý Halal Nội bộ"),
]


def build(doc: DocxDocument, content: str, title: str, filename: str, cfg: dict = {}):
    conf_label = cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ")
    B.setup_page(doc, title, confidential_label=conf_label)

    # Cover metadata — use admin config or defaults
    raw_meta = cfg.get("cover_meta")
    if raw_meta:
        extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")]
    else:
        extra_meta = _DEFAULT_META

    show_toc = cfg.get("show_toc", True)
    show_rev = cfg.get("show_revision_table", True)
    show_appr = cfg.get("show_approval_block", True)

    B.add_cover(doc,
        doc_type_label="Halal Policy",
        title=title,
        extra_meta=extra_meta,
        show_approval=show_appr,
    )

    # Table of contents
    if show_toc:
        B.add_heading(doc, "MỤC LỤC", level=1)
        toc_p = doc.add_paragraph()
        r = toc_p.add_run("[Cập nhật mục lục: References → Update Table]")
        B.font(r, size=B.SZ_SMALL, italic=True, color=B.CLR_GRAY)
        fld = (
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' w:instr=" TOC \\o &quot;1-3&quot; \\h \\z \\u ">'
            '<w:r><w:t></w:t></w:r></w:fldSimple>'
        )
        toc_p._element.append(etree.fromstring(fld))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Revision history
    if show_rev:
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Custom sections from admin
    for sec in cfg.get("custom_sections", []):
        if sec.get("title"):
            B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "):
                    B.add_bullet(doc, line[2:])
                elif line:
                    B.add_body(doc, line)

    # Main content
    B.parse_content(doc, content)
    B.add_ending(doc)
