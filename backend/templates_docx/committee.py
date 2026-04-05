"""Template: Internal Halal Committee."""
from docx import Document as DocxDocument
from docx.shared import Cm
from docx.enum.text import WD_BREAK
from . import _base as B

def build(doc: DocxDocument, content: str, title: str, filename: str, cfg: dict = {}):
    B.setup_page(doc, title, confidential_label=cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ"))
    raw_meta = cfg.get("cover_meta")
    extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")] if raw_meta else [
        ("Tiêu chuẩn", "MS 1500:2019 Clause 3.5"),
        ("Cơ quan phê duyệt", "Ban Giám đốc"),
    ]
    B.add_cover(doc, doc_type_label="Internal Halal Committee", title=title,
                extra_meta=extra_meta, show_approval=cfg.get("show_approval_block", True))
    if cfg.get("show_revision_table", True):
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    B.add_heading(doc, "CƠ CẤU BAN QUẢN LÝ HALAL NỘI BỘ", level=1)
    tbl = doc.add_table(rows=5, cols=4)
    tbl.style = "Table Grid"
    for ci, h in enumerate(["STT", "Vị trí", "Họ tên", "Bộ phận"]):
        cell = tbl.cell(0, ci); cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        B.font(r, size=B.SZ_TINY, bold=True, color=B.CLR_WHITE)
        cell.paragraphs[0].alignment = 1
        B.shade_cell(cell, "0D47A1"); B.cell_padding(cell)
    for ri, (stt, role) in enumerate([("1","Trưởng ban"),("2","Phó ban"),("3","Thư ký"),("4","Thành viên")], 1):
        for ci, val in enumerate((stt, role, "[Họ tên]", "[Bộ phận]")):
            cell = tbl.cell(ri, ci); cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            is_p = val.startswith("[")
            B.font(r, size=B.SZ_SMALL, italic=is_p, color=B.CLR_GRAY if is_p else B.CLR_BLACK)
            B.cell_padding(cell)
    doc.add_paragraph()
    for sec in cfg.get("custom_sections", []):
        if sec.get("title"): B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "): B.add_bullet(doc, line[2:])
                elif line: B.add_body(doc, line)
    B.parse_content(doc, content)
    B.add_ending(doc)
