"""Template: SOP — Standard Operating Procedure (dùng chung cho 6 loại SOP)."""
from docx import Document as DocxDocument
from docx.shared import Cm
from docx.enum.text import WD_BREAK
from . import _base as B

SOP_NAMES = {
    "sop_raw_material_receiving":   "Tiếp nhận nguyên liệu thô",
    "sop_storage_segregation":      "Lưu kho và phân tách",
    "sop_production_operation":     "Vận hành sản xuất",
    "sop_cleaning_sanitation":      "Vệ sinh và khử trùng",
    "sop_handling_nonconformances": "Xử lý sự không phù hợp",
    "sop_complaint_recall":         "Khiếu nại và thu hồi sản phẩm",
}


def build(doc: DocxDocument, content: str, title: str, filename: str,
          doc_type: str = "sop_production_operation", cfg: dict = {}):
    sop_name = SOP_NAMES.get(doc_type, "Quy trình vận hành")
    conf_label = cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ")

    B.setup_page(doc, f"SOP — {sop_name}", confidential_label=conf_label)

    raw_meta = cfg.get("cover_meta")
    extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")] if raw_meta else [
        ("Loại SOP", sop_name),
        ("Tiêu chuẩn", "MS 1500:2019, GMP, HACCP"),
        ("Tần suất rà soát", "6 tháng / lần"),
    ]

    B.add_cover(doc,
        doc_type_label=f"SOP — {sop_name}",
        title=title,
        extra_meta=extra_meta,
        show_approval=cfg.get("show_approval_block", True),
    )

    if cfg.get("show_revision_table", True):
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # SOP info table
    B.add_heading(doc, "THÔNG TIN QUY TRÌNH", level=1)
    info_tbl = doc.add_table(rows=5, cols=2)
    info_tbl.style = "Table Grid"
    info_rows = [
        ("Tên quy trình", sop_name),
        ("Mã quy trình", f"SOP-{doc_type.split('_')[-1][:4].upper()}-001"),
        ("Bộ phận thực hiện", "[Điền tên bộ phận]"),
        ("Người chịu trách nhiệm", "[Điền họ tên]"),
        ("Tần suất thực hiện", "[Hàng ngày / Hàng tuần / Theo lô]"),
    ]
    for ri, (label, value) in enumerate(info_rows):
        c0 = info_tbl.cell(ri, 0); c1 = info_tbl.cell(ri, 1)
        c0.width = Cm(5); c1.width = Cm(9)
        c0.text = ""; c1.text = ""
        r0 = c0.paragraphs[0].add_run(label)
        B.font(r0, size=B.SZ_SMALL, bold=True)
        r1 = c1.paragraphs[0].add_run(value)
        B.font(r1, size=B.SZ_SMALL)
        B.shade_cell(c0, "F5F5F5")
        B.cell_padding(c0); B.cell_padding(c1)
    doc.add_paragraph()

    for sec in cfg.get("custom_sections", []):
        if sec.get("title"):
            B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "):
                    B.add_bullet(doc, line[2:])
                elif line:
                    B.add_body(doc, line)

    B.parse_content(doc, content)

    doc.add_paragraph()
    B.add_heading(doc, "BIỂU MẪU KIỂM SOÁT", level=1)
    note_p = doc.add_paragraph()
    r = note_p.add_run("[Đính kèm biểu mẫu kiểm soát / checklist tương ứng với quy trình này]")
    B.font(r, size=B.SZ_SMALL, italic=True, color=B.CLR_GRAY)

    B.add_ending(doc)
