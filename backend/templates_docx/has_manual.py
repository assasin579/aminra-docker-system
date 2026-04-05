"""Template: HAS Manual — Halal Assurance System Manual."""
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from lxml import etree
from . import _base as B

_DEFAULT_META = [
    ("Tiêu chuẩn", "HAS 23000 / MS 1500:2019"),
    ("Phạm vi", "Toàn bộ hệ thống đảm bảo Halal"),
    ("Bộ phận", "Ban Quản lý Halal Nội bộ"),
]


def build(doc: DocxDocument, content: str, title: str, filename: str, cfg: dict = {}):
    conf_label = cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ")
    B.setup_page(doc, title, confidential_label=conf_label)

    raw_meta = cfg.get("cover_meta")
    extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")] if raw_meta else _DEFAULT_META

    B.add_cover(doc,
        doc_type_label="HAS Manual",
        title=title,
        extra_meta=extra_meta,
        show_approval=cfg.get("show_approval_block", True),
    )

    if cfg.get("show_toc", True):
        B.add_heading(doc, "MỤC LỤC", level=1)
        toc_p = doc.add_paragraph()
        r = toc_p.add_run("[Cập nhật mục lục: References → Update Table]")
        B.font(r, size=B.SZ_SMALL, italic=True, color=B.CLR_GRAY)
        fld = (
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' w:instr=" TOC \\o &quot;1-3&quot; \\h \\z \\u ">'
            '<w:r><w:t></w:t></w:r></w:fldSimple>'
        )
        toc_p._element.append(etree.fromstring(fld))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if cfg.get("show_revision_table", True):
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for sec in cfg.get("custom_sections", []):
        if sec.get("title"):
            B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "):
                    B.add_bullet(doc, line[2:])
                elif line:
                    B.add_body(doc, line)

    B.parse_content(doc, content)
    B.add_ending(doc)
