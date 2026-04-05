"""Template: Ingredient & Raw Material Documentation."""
from docx import Document as DocxDocument
from docx.shared import Cm
from docx.enum.text import WD_BREAK
from . import _base as B

def build(doc: DocxDocument, content: str, title: str, filename: str, cfg: dict = {}):
    conf_label = cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ")
    B.setup_page(doc, title, confidential_label=conf_label)
    raw_meta = cfg.get("cover_meta")
    extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")] if raw_meta else [
        ("Tiêu chuẩn", "MS 1500:2019, Codex Alimentarius"),
        ("Bộ phận", "QA/QC & Mua hàng (Procurement)"),
        ("Tần suất cập nhật", "Mỗi khi thay đổi nhà cung cấp / nguyên liệu"),
    ]
    B.add_cover(doc, doc_type_label="Ingredient & Raw Material", title=title,
                extra_meta=extra_meta, show_approval=cfg.get("show_approval_block", True))
    if cfg.get("show_revision_table", True):
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    B.add_heading(doc, "BẢNG KIỂM TRA NHÀ CUNG CẤP", level=1)
    chk_tbl = doc.add_table(rows=2, cols=4)
    chk_tbl.style = "Table Grid"
    for ci, h in enumerate(["Nhà cung cấp", "Nguyên liệu", "Chứng nhận Halal", "Hạn hiệu lực"]):
        cell = chk_tbl.cell(0, ci); cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        B.font(r, size=B.SZ_TINY, bold=True, color=B.CLR_WHITE)
        B.shade_cell(cell, "0D47A1"); B.cell_padding(cell)
    for ci in range(4):
        cell = chk_tbl.cell(1, ci); cell.text = ""
        r = cell.paragraphs[0].add_run("[Điền thông tin]")
        B.font(r, size=B.SZ_TINY, italic=True, color=B.CLR_GRAY); B.cell_padding(cell)
    doc.add_paragraph()
    for sec in cfg.get("custom_sections", []):
        if sec.get("title"): B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "): B.add_bullet(doc, line[2:])
                elif line: B.add_body(doc, line)
    B.parse_content(doc, content)
    B.add_ending(doc)
