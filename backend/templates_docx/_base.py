"""
Shared DOCX building blocks — cover page, header/footer, styles, tables, parsing.
Every template imports from here; no template duplicates this logic.
"""
from __future__ import annotations
import re
from datetime import datetime
from docx import Document as DocxDocument
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from lxml import etree

# ── Constants ─────────────────────────────────────────────────────────────────

FONT       = "Times New Roman"
SZ_BODY    = Pt(13)
SZ_SMALL   = Pt(11)
SZ_TINY    = Pt(9)
SZ_H1      = Pt(16)
SZ_H2      = Pt(14)
SZ_H3      = Pt(13)
SZ_TITLE   = Pt(24)
CLR_BLACK  = RGBColor(0x1A, 0x1A, 0x1A)
CLR_GRAY   = RGBColor(0x66, 0x66, 0x66)
CLR_LIGHT  = RGBColor(0x99, 0x99, 0x99)
CLR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CLR_ACCENT = RGBColor(0x0D, 0x47, 0xA1)
LINE_SP    = 1.5

# ── Low-level helpers ─────────────────────────────────────────────────────────

def font(run, size=SZ_BODY, bold=False, italic=False, color=CLR_BLACK, name=FONT):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), name)

def spacing(p, before=Pt(0), after=Pt(6), line=LINE_SP):
    pf = p.paragraph_format
    pf.space_before = before
    pf.space_after  = after
    pf.line_spacing = line

def hr(p, color="1A1A1A", sz="8"):
    bdr = p._element.get_or_add_pPr().makeelement(qn("w:pBdr"), {})
    bdr.append(bdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): sz,
        qn("w:space"): "1", qn("w:color"): color,
    }))
    p._element.get_or_add_pPr().append(bdr)

def shade_cell(cell, fill):
    tc = cell._element.get_or_add_tcPr()
    tc.append(tc.makeelement(qn("w:shd"), {
        qn("w:fill"): fill, qn("w:val"): "clear",
    }))

def cell_padding(cell, top=60, bottom=60):
    tc = cell._element.get_or_add_tcPr()
    mar = tc.makeelement(qn("w:tcMar"), {})
    for side, val in (("top", str(top)), ("bottom", str(bottom))):
        mar.append(mar.makeelement(qn(f"w:{side}"), {
            qn("w:w"): val, qn("w:type"): "dxa",
        }))
    tc.append(mar)

def inline(p, text, base_size=SZ_BODY, base_bold=False, base_color=CLR_BLACK):
    """Parse **bold** and *italic* within text."""
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            font(r, size=base_size, bold=True, color=base_color)
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            font(r, size=base_size, italic=True, color=base_color)
        else:
            r = p.add_run(part)
            font(r, size=base_size, bold=base_bold, color=base_color)


# ── Page setup ────────────────────────────────────────────────────────────────

def setup_page(doc: DocxDocument, title: str, confidential_label="TÀI LIỆU NỘI BỘ"):
    """Configure margins, header with doc title, footer with page number."""
    for section in doc.sections:
        section.top_margin    = Cm(2.54)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

        # ── Header ──
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr_run = hp.add_run(title)
        font(hr_run, size=Pt(8), italic=True, color=CLR_LIGHT)
        hp.add_run("    ")
        conf_run = hp.add_run(confidential_label)
        font(conf_run, size=Pt(8), bold=True, color=CLR_ACCENT)
        # Header bottom border
        pPr = hp._element.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        pBdr.append(pBdr.makeelement(qn("w:bottom"), {
            qn("w:val"): "single", qn("w:sz"): "4",
            qn("w:space"): "4", qn("w:color"): "CCCCCC",
        }))
        pPr.append(pBdr)

        # ── Footer ──
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Trang ")
        font(fr, size=SZ_TINY, color=CLR_LIGHT)
        fld = (
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' w:instr=" PAGE "><w:r><w:rPr><w:sz w:val="18"/><w:color w:val="999999"/>'
            '</w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        )
        fp._element.append(etree.fromstring(fld))


# ── Cover page ────────────────────────────────────────────────────────────────

def add_cover(doc: DocxDocument, *,
              doc_type_label: str,
              title: str,
              version: str = "1.0",
              extra_meta: list[tuple[str, str]] | None = None,
              show_approval: bool = True):
    """Professional cover page — title, metadata table, approval block."""
    today = datetime.now().strftime("%d/%m/%Y")

    # Spacer
    for _ in range(4):
        sp = doc.add_paragraph(); spacing(sp, after=Pt(0))

    # Document type badge
    badge_p = doc.add_paragraph()
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = badge_p.add_run(f"  {doc_type_label.upper()}  ")
    font(br, size=SZ_SMALL, bold=True, color=CLR_ACCENT)
    spacing(badge_p, after=Pt(16))

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run(title.upper())
    font(tr, size=SZ_TITLE, bold=True, color=CLR_BLACK)
    spacing(title_p, after=Pt(8))

    # Decorative line
    line_p = doc.add_paragraph()
    line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr(line_p, color="0D47A1", sz="12")
    spacing(line_p, after=Pt(24))

    # Metadata table (borderless)
    meta_rows = [
        ("Mã tài liệu", f"AMINRA-{doc_type_label[:3].upper()}-{version.replace('.', '')}"),
        ("Phiên bản", f"v{version}"),
        ("Ngày ban hành", today),
        ("Phân loại", "Nội bộ — Hạn chế"),
    ]
    if extra_meta:
        meta_rows.extend(extra_meta)

    tbl = doc.add_table(rows=len(meta_rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, (label, value) in enumerate(meta_rows):
        c0 = tbl.cell(ri, 0); c1 = tbl.cell(ri, 1)
        c0.width = Cm(5); c1.width = Cm(8)
        c0.text = ""; c1.text = ""
        r0 = c0.paragraphs[0].add_run(label)
        font(r0, size=SZ_SMALL, bold=True, color=CLR_GRAY)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = c1.paragraphs[0].add_run(value)
        font(r1, size=SZ_SMALL, color=CLR_BLACK)
        for c in (c0, c1):
            cell_padding(c, top=40, bottom=40)
    # Remove borders from metadata table
    tbl_xml = tbl._element
    tblPr = tbl_xml.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = tbl_xml.makeelement(qn("w:tblPr"), {})
        tbl_xml.insert(0, tblPr)
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(borders.makeelement(qn(f"w:{side}"), {
            qn("w:val"): "none", qn("w:sz"): "0",
            qn("w:space"): "0", qn("w:color"): "auto",
        }))
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(borders)

    # Spacer then approval block
    if show_approval:
        doc.add_paragraph()
        _add_approval_block(doc)

    # Page break
    br_p = doc.add_paragraph()
    br_p.add_run().add_break(WD_BREAK.PAGE)


def _add_approval_block(doc: DocxDocument):
    """Three-column approval table: Soạn thảo / Xem xét / Phê duyệt."""
    tbl = doc.add_table(rows=4, cols=3)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Soạn thảo", "Xem xét", "Phê duyệt"]
    fields  = ["Họ tên:", "Chức vụ:", "Chữ ký / Ngày:"]

    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        font(r, size=SZ_SMALL, bold=True, color=CLR_WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, "0D47A1")
        cell_padding(cell, top=50, bottom=50)

    for ri, field in enumerate(fields, start=1):
        for ci in range(3):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            r = cell.paragraphs[0].add_run(field)
            font(r, size=SZ_TINY, color=CLR_GRAY)
            cell.height = Cm(1.5)
            cell_padding(cell, top=40, bottom=40)


# ── Revision history table ────────────────────────────────────────────────────

def add_revision_table(doc: DocxDocument, title="LỊCH SỬ THAY ĐỔI TÀI LIỆU"):
    add_heading(doc, title, level=1)
    tbl = doc.add_table(rows=4, cols=5)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads = ["Phiên bản", "Ngày", "Nội dung thay đổi", "Người thực hiện", "Phê duyệt"]
    for ci, h in enumerate(heads):
        cell = tbl.cell(0, ci)
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        font(r, size=SZ_TINY, bold=True, color=CLR_WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, "0D47A1")
        cell_padding(cell, top=40, bottom=40)

    today = datetime.now().strftime("%d/%m/%Y")
    row1 = ["1.0", today, "Phiên bản đầu tiên (AI-assisted)", "AMINRA AI", ""]
    for ci, val in enumerate(row1):
        cell = tbl.cell(1, ci)
        cell.text = ""
        r = cell.paragraphs[0].add_run(val)
        font(r, size=SZ_TINY, color=CLR_BLACK)
        cell_padding(cell, top=30, bottom=30)

    # Empty rows for future revisions
    for ri in (2, 3):
        for ci in range(5):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            cell_padding(cell, top=30, bottom=30)

    doc.add_paragraph()


# ── Content helpers ───────────────────────────────────────────────────────────

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level <= 1:
        font(run, size=SZ_H1, bold=True, color=CLR_BLACK)
        spacing(p, before=Pt(24), after=Pt(8), line=1.2)
        hr(p, color="1A1A1A", sz="6")
    elif level == 2:
        font(run, size=SZ_H2, bold=True, color=CLR_BLACK)
        spacing(p, before=Pt(18), after=Pt(6), line=1.2)
    else:
        font(run, size=SZ_H3, bold=True, italic=True, color=CLR_BLACK)
        spacing(p, before=Pt(12), after=Pt(4), line=1.2)
    return p

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    inline(p, text)
    spacing(p, after=Pt(6))
    return p

def add_bullet(doc, text, indent=0):
    p = doc.add_paragraph(style="List Bullet")
    if indent > 0:
        p.paragraph_format.left_indent = Cm(1.27 * (indent + 1))
    p.clear()
    inline(p, text)
    spacing(p, after=Pt(3))
    return p

def add_numbered_item(doc, text, indent=0):
    p = doc.add_paragraph(style="List Number")
    if indent > 0:
        p.paragraph_format.left_indent = Cm(1.27 * (indent + 1))
    p.clear()
    inline(p, text)
    spacing(p, after=Pt(3))
    return p

def add_table_from_lines(doc, rows_text):
    sep = "|" if "|" in rows_text[0] else "\t"
    parsed = []
    for row in rows_text:
        cells = [c.strip() for c in row.split(sep)]
        cells = [c for c in cells if c]
        if cells and not all(set(c) <= {"-", ":", " ", "="} for c in cells):
            parsed.append(cells)
    if not parsed:
        return
    cols = max(len(r) for r in parsed)
    tbl = doc.add_table(rows=len(parsed), cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, cells in enumerate(parsed):
        for ci, txt in enumerate(cells):
            if ci >= cols: break
            cell = tbl.cell(ri, ci)
            cell.text = ""
            r = cell.paragraphs[0].add_run(txt)
            is_hdr = ri == 0
            font(r, size=SZ_SMALL, bold=is_hdr, color=CLR_BLACK)
            spacing(cell.paragraphs[0], before=Pt(3), after=Pt(3), line=1.15)
            cell_padding(cell)
        for ci in range(len(cells), cols):
            tbl.cell(ri, ci).text = ""
    # Header shading
    for cell in tbl.rows[0].cells:
        shade_cell(cell, "E8E8E8")
    doc.add_paragraph()


# ── Document ending ───────────────────────────────────────────────────────────

def add_ending(doc: DocxDocument):
    today = datetime.now().strftime("%d/%m/%Y")
    doc.add_paragraph()
    end_p = doc.add_paragraph()
    hr(end_p, color="999999", sz="4")
    spacing(end_p, before=Pt(18))

    sig_p = doc.add_paragraph()
    sig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sig_p.add_run("— HẾT —")
    font(sr, size=SZ_BODY, bold=True, color=CLR_GRAY)
    spacing(sig_p, before=Pt(6), after=Pt(12))

    note_p = doc.add_paragraph()
    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note_p.add_run(
        f"Tạo ngày {today} bởi AMINRA — Halal Certification AI Platform\n"
        "Vui lòng rà soát và phê duyệt trước khi ban hành chính thức."
    )
    font(nr, size=SZ_TINY, italic=True, color=CLR_LIGHT)


# ── Universal content parser ──────────────────────────────────────────────────

_RE_NUM     = re.compile(r'^([\d]+(?:\.[\d]+)*)[.\s:]+\s*(.*)')
_RE_MD_H1   = re.compile(r'^#{1}\s+(.+)')
_RE_MD_H2   = re.compile(r'^#{2}\s+(.+)')
_RE_MD_H3   = re.compile(r'^#{3,}\s+(.+)')
_RE_LETTER  = re.compile(r'^[a-z]\)\s+(.+)')
_RE_ROMAN   = re.compile(r'^[ivxIVX]+[.)]\s+(.+)')
_RE_NUMLIST = re.compile(r'^(\d{1,2})[.)]\s+(.+)')


def _is_heading_text(rest: str) -> bool:
    if not rest:
        return True
    if len(rest) > 100:
        return False
    if rest.rstrip().endswith(('.', ',', ';')):
        return False
    if rest.count('. ') > 1:
        return False
    words = rest.split()
    if len(words) <= 12:
        upper_count = sum(1 for w in words if w[0:1].isupper() or not w[0:1].isalpha())
        if upper_count >= len(words) * 0.5:
            return True
    return len(rest) < 60


def parse_content(doc: DocxDocument, content: str):
    """Parse LLM-generated text and render into the document."""
    lines = content.split("\n")
    idx = 0
    table_buf: list[str] = []
    in_table = False
    empty_count = 0

    while idx < len(lines):
        raw = lines[idx]
        s = raw.strip()

        # Table detection
        is_tbl = ("|" in s and s.count("|") >= 2) or ("\t" in raw and raw.count("\t") >= 1)
        if is_tbl:
            table_buf.append(s)
            in_table = True
            idx += 1
            continue
        elif in_table and table_buf:
            add_table_from_lines(doc, table_buf)
            table_buf = []
            in_table = False

        if not s:
            empty_count += 1
            if empty_count <= 2:
                sp = doc.add_paragraph(); spacing(sp, after=Pt(2))
            idx += 1
            continue
        empty_count = 0

        # Markdown headings
        m = _RE_MD_H1.match(s)
        if m: add_heading(doc, m.group(1).strip(), 1); idx += 1; continue
        m = _RE_MD_H2.match(s)
        if m: add_heading(doc, m.group(1).strip(), 2); idx += 1; continue
        m = _RE_MD_H3.match(s)
        if m: add_heading(doc, m.group(1).strip(), 3); idx += 1; continue

        # Numbered section heading
        m = _RE_NUM.match(s)
        if m:
            num_part = m.group(1)
            rest = m.group(2).strip()
            if _is_heading_text(rest):
                depth = num_part.count(".")
                add_heading(doc, s, level=min(depth + 1, 3))
                idx += 1
                continue

        # ALL-CAPS heading
        if (s == s.upper() and 3 < len(s) < 80
                and s[0].isalpha() and not s.endswith(('.', ',', ';'))):
            add_heading(doc, s, 1)
            idx += 1
            continue

        # Letter sub-item
        m = _RE_LETTER.match(s)
        if m: add_bullet(doc, m.group(1), indent=1); idx += 1; continue

        # Roman numeral
        m = _RE_ROMAN.match(s)
        if m: add_bullet(doc, m.group(1), indent=1); idx += 1; continue

        # Bullet
        if s[:2] in ("- ", "* ") or s.startswith("• "):
            add_bullet(doc, s[2:].strip()); idx += 1; continue

        # Numbered list item
        m = _RE_NUMLIST.match(s)
        if m and not re.match(r'^\d+\.\d+', s):
            add_numbered_item(doc, m.group(2)); idx += 1; continue

        # Body paragraph
        add_body(doc, s)
        idx += 1

    if table_buf:
        add_table_from_lines(doc, table_buf)
