"""Template: Process Flow Chart."""
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from . import _base as B

def build(doc: DocxDocument, content: str, title: str, filename: str, cfg: dict = {}):
    B.setup_page(doc, title, confidential_label=cfg.get("confidential_label", "TÀI LIỆU NỘI BỘ"))
    raw_meta = cfg.get("cover_meta")
    extra_meta = [(m["key"], m["value"]) for m in raw_meta if m.get("key")] if raw_meta else [
        ("Tiêu chuẩn", "MS 1500:2019, HACCP Codex"),
        ("Bộ phận", "Sản xuất & QA/QC"),
    ]
    B.add_cover(doc, doc_type_label="Process Flow Chart", title=title,
                extra_meta=extra_meta, show_approval=cfg.get("show_approval_block", True))
    if cfg.get("show_revision_table", True):
        B.add_revision_table(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    B.add_heading(doc, "SƠ ĐỒ QUY TRÌNH TỔNG QUÁT", level=1)
    box_p = doc.add_paragraph(); box_p.alignment = 1
    r = box_p.add_run("[Chèn sơ đồ quy trình tại đây — Insert → SmartArt → Process]")
    B.font(r, size=B.SZ_SMALL, italic=True, color=B.CLR_GRAY)
    doc.add_paragraph()
    for sec in cfg.get("custom_sections", []):
        if sec.get("title"): B.add_heading(doc, sec["title"], level=1)
        if sec.get("content"):
            for line in sec["content"].split("\n"):
                line = line.strip()
                if line.startswith("- "): B.add_bullet(doc, line[2:])
                elif line: B.add_body(doc, line)
    B.parse_content(doc, content)
    B.add_ending(doc)
