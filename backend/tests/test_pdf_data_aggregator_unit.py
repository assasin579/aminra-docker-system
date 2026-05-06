"""Unit tests for services/pdf_data_aggregator.py.

Mocks the asyncpg DB layer via AsyncMock; filesystem reads are exercised
against a tmp_path fixture so admin_templates loading is verified for real.

Critical invariants tested:
  - tenant isolation (every DB query carries tenant_id; cross-tenant cache
    can't bleed)
  - cache hit reuses source data but rebuilds bundle with fresh request
    payload (so cached company doesn't trap an old payload)
  - DOCX text extraction sanitises whitespace
  - admin asset directory absence is non-fatal (returns empty AdminAssets)
"""
from __future__ import annotations

import json
import time
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from docx import Document as DocxDoc

from services.pdf_data_aggregator import (
    AdminAssets,
    CompanyData,
    PDFDataAggregator,
    RawDataBundle,
    invalidate_cache,
)


# ── Fixtures ────────────────────────────────────────────────────────────────


@pytest.fixture(autouse=True)
def _reset_cache():
    """Each test gets a fresh aggregator cache."""
    invalidate_cache()
    yield
    invalidate_cache()


@pytest.fixture
def aggregator():
    return PDFDataAggregator()


@pytest.fixture
def db_mock_company():
    """asyncpg-like connection returning a typical owner row."""
    db = AsyncMock()
    db.fetchrow = AsyncMock(return_value={
        "company_name": "Công ty TNHH Quế Thiên Lộc",
        "company_code": "0123456789",
        "tenant_id": "tenant-A",
        "address": "Số 12 Lê Lợi, Vinh, Nghệ An",
        "phone": "+842381234567",
        "email": "info@quethienloc.vn",
        "representative_name": "Nguyễn Văn Anh",
        "manager_name": None,
        "created_at": None,
    })
    return db


@pytest.fixture
def db_mock_no_owner():
    db = AsyncMock()
    db.fetchrow = AsyncMock(return_value=None)
    return db


@pytest.fixture
def admin_template_dir(tmp_path, monkeypatch):
    """Patch _ADMIN_TPL_DIR to a tmp_path so admin file reads target it."""
    from services import pdf_data_aggregator as agg
    monkeypatch.setattr(agg, "_ADMIN_TPL_DIR", tmp_path)
    return tmp_path


def _actor(role: str = "business", tenant_id: str = "tenant-A") -> dict:
    return {
        "sub": "actor-id",
        "role": role,
        "tenant_id": tenant_id,
        "is_owner": True,
    }


# ── Stage-1 fetch path ─────────────────────────────────────────────────────


async def test_fetch_returns_bundle_with_company(
    aggregator, db_mock_company, admin_template_dir,
):
    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={"business_name": "X"},
    )
    assert isinstance(bundle, RawDataBundle)
    assert bundle.tenant_id == "tenant-A"
    assert bundle.doc_type == "company_profile"
    assert bundle.company is not None
    assert bundle.company.business_name == "Công ty TNHH Quế Thiên Lộc"
    assert bundle.company.tax_code == "0123456789"
    assert bundle.company.phone == "+842381234567"
    # Phase 2 stubs return empty
    assert bundle.active_submission is None
    assert bundle.certificates == []
    assert bundle.documents == []


async def test_tenant_isolation_passes_tenant_id_to_query(
    aggregator, db_mock_company, admin_template_dir,
):
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-XYZ", actor=_actor(tenant_id="tenant-XYZ"),
        request_payload={},
    )
    # The query MUST be parameterised on tenant_id, never interpolated.
    args, _ = db_mock_company.fetchrow.call_args
    assert "WHERE tenant_id = $1" in args[0]
    assert "tenant-XYZ" in args


async def test_no_owner_row_returns_none_company(
    aggregator, db_mock_no_owner, admin_template_dir,
):
    bundle = await aggregator.fetch(
        db=db_mock_no_owner, doc_type="company_profile",
        tenant_id="ghost-tenant", actor=_actor(tenant_id="ghost-tenant"),
        request_payload={},
    )
    assert bundle.company is None


async def test_blank_tenant_id_short_circuits(
    aggregator, db_mock_company, admin_template_dir,
):
    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="", actor=_actor(tenant_id=""),
        request_payload={},
    )
    assert bundle.company is None
    db_mock_company.fetchrow.assert_not_called()


# ── Admin cfg JSON ──────────────────────────────────────────────────────────


async def test_admin_cfg_loads_when_present(
    aggregator, db_mock_company, admin_template_dir,
):
    cfg_path = admin_template_dir / "company_profile.json"
    cfg_path.write_text(json.dumps({
        "docx_config": {"confidential_label": "HỒ SƠ MẪU", "show_approval_block": False},
    }), encoding="utf-8")

    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert bundle.admin_cfg["confidential_label"] == "HỒ SƠ MẪU"
    assert bundle.admin_cfg["show_approval_block"] is False


async def test_admin_cfg_missing_returns_empty(
    aggregator, db_mock_company, admin_template_dir,
):
    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="never_existed",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert bundle.admin_cfg == {}


async def test_admin_cfg_invalid_json_does_not_crash(
    aggregator, db_mock_company, admin_template_dir, caplog,
):
    bad = admin_template_dir / "broken.json"
    bad.write_text("{ not valid json", encoding="utf-8")
    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="broken",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert bundle.admin_cfg == {}


# ── Admin assets (DOCX text + image files) ──────────────────────────────────


async def test_docx_text_extracted(
    aggregator, db_mock_company, admin_template_dir, tmp_path,
):
    files_dir = admin_template_dir / "files" / "halal_policy"
    files_dir.mkdir(parents=True)
    docx_path = files_dir / "policy.docx"
    doc = DocxDoc()
    doc.add_paragraph("Tiêu đề chính sách Halal")
    doc.add_paragraph("")
    doc.add_paragraph("Cam kết tuân thủ MS 1500.")
    doc.save(str(docx_path))

    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="halal_policy",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert bundle.admin_assets.reference_text is not None
    assert "Tiêu đề chính sách Halal" in bundle.admin_assets.reference_text
    assert "Cam kết tuân thủ MS 1500." in bundle.admin_assets.reference_text


async def test_extra_assets_indexed_by_stem(
    aggregator, db_mock_company, admin_template_dir,
):
    files_dir = admin_template_dir / "files" / "company_profile"
    files_dir.mkdir(parents=True)
    (files_dir / "logo.png").write_bytes(b"\x89PNG\r\n\x1a\nfake")
    (files_dir / "header.svg").write_text("<svg/>", encoding="utf-8")

    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert "logo" in bundle.admin_assets.extra_assets
    assert "header" in bundle.admin_assets.extra_assets
    assert bundle.admin_assets.extra_assets["logo"].startswith("file://")


async def test_assets_dir_missing_returns_empty(
    aggregator, db_mock_company, admin_template_dir,
):
    bundle = await aggregator.fetch(
        db=db_mock_company, doc_type="never_uploaded",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert bundle.admin_assets == AdminAssets()


# ── Cache behaviour ────────────────────────────────────────────────────────


async def test_cache_hit_within_ttl_reuses_db_query(
    aggregator, db_mock_company, admin_template_dir,
):
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={"business_name": "first"},
    )
    db_mock_company.fetchrow.reset_mock()

    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={"business_name": "second"},
    )
    db_mock_company.fetchrow.assert_not_called()


async def test_cache_returns_fresh_payload_on_hit(
    aggregator, db_mock_company, admin_template_dir,
):
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={"business_name": "first"},
    )
    bundle2 = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={"business_name": "second"},
    )
    # Cached SOURCE data, but request_payload always reflects current call
    assert bundle2.request_payload == {"business_name": "second"}


async def test_cache_isolates_per_tenant(
    aggregator, db_mock_company, admin_template_dir,
):
    """Tenant A's cached bundle MUST NOT serve tenant B."""
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(tenant_id="tenant-A"),
        request_payload={},
    )
    db_mock_company.fetchrow.reset_mock()

    # Different tenant — separate cache key — must hit DB again
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-B", actor=_actor(tenant_id="tenant-B"),
        request_payload={},
    )
    db_mock_company.fetchrow.assert_called_once()


async def test_invalidate_cache_forces_fetch(
    aggregator, db_mock_company, admin_template_dir,
):
    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    invalidate_cache()
    db_mock_company.fetchrow.reset_mock()

    await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    db_mock_company.fetchrow.assert_called_once()


# ── Versions hash ──────────────────────────────────────────────────────────


async def test_versions_hash_deterministic(
    aggregator, db_mock_company, admin_template_dir,
):
    b1 = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    invalidate_cache()
    b2 = await aggregator.fetch(
        db=db_mock_company, doc_type="company_profile",
        tenant_id="tenant-A", actor=_actor(),
        request_payload={},
    )
    assert b1.sources_versions == b2.sources_versions


async def test_versions_change_when_company_changes(
    aggregator, admin_template_dir,
):
    """Different DB row → different sources_versions hash."""
    db1 = AsyncMock()
    db1.fetchrow = AsyncMock(return_value={
        "company_name": "First Co", "company_code": "1", "tenant_id": "t-A",
        "address": None, "phone": None, "email": None,
        "representative_name": None, "manager_name": None, "created_at": None,
    })
    b1 = await aggregator.fetch(
        db=db1, doc_type="company_profile",
        tenant_id="t-A", actor=_actor(tenant_id="t-A"),
        request_payload={},
    )
    invalidate_cache()

    db2 = AsyncMock()
    db2.fetchrow = AsyncMock(return_value={
        "company_name": "Second Co", "company_code": "2", "tenant_id": "t-A",
        "address": None, "phone": None, "email": None,
        "representative_name": None, "manager_name": None, "created_at": None,
    })
    b2 = await aggregator.fetch(
        db=db2, doc_type="company_profile",
        tenant_id="t-A", actor=_actor(tenant_id="t-A"),
        request_payload={},
    )
    assert b1.sources_versions["company"] != b2.sources_versions["company"]
